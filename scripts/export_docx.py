#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Enterprise Word (.docx) Documentation Exporter
Generates comprehensive Chinese (zh) and Japanese (ja) Word documents
from Markdown files in docs/zh and docs/ja.
"""

import os
import re
import sys
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path("/home/fx/Develop/SnowflakeBlog")
EXPORTS_DIR = BASE_DIR / "exports"
EXPORTS_DIR.mkdir(exist_ok=True)

CHAPTERS = [
    ("00_architecture_overview.md", "00"),
    ("01_snowflake_signup.md", "01"),
    ("02_oracle_setup_and_data.md", "02"),
    ("03_snowflake_rbac_setup.md", "03"),
    ("04_data_migration.md", "04"),
    ("05_data_reconciliation.md", "05"),
    ("06_troubleshooting_tips.md", "06"),
]

# Color Palette (Enterprise Snowflake Theme)
COLOR_PRIMARY = RGBColor(11, 60, 93)      # Deep Navy #0B3C5D
COLOR_SECONDARY = RGBColor(37, 99, 235)   # Tech Blue #2563EB
COLOR_TEXT_MAIN = RGBColor(30, 41, 59)    # Slate Dark #1E293B
COLOR_TEXT_MUTED = RGBColor(100, 116, 139) # Slate Muted #64748B
COLOR_CODE_TEXT = RGBColor(15, 23, 42)    # #0F172A
COLOR_BORDER = "CBD5E1"
COLOR_HEADER_BG = "1E3A8A"                # Deep Royal Blue
COLOR_ZEBRA_BG = "F8FAFC"                 # Subtle Slate
COLOR_CODE_BG = "F1F5F9"                  # Code block background

CALLOUT_CONFIGS = {
    "TIP": {
        "border_color": "059669",
        "bg_color": "ECFDF5",
        "title_zh": "💡 专家建议与最佳实践 (TIP)",
        "title_ja": "💡 ベストプラクティス・推奨方針 (TIP)",
        "title_color": RGBColor(5, 150, 105),
    },
    "NOTE": {
        "border_color": "2563EB",
        "bg_color": "EFF6FF",
        "title_zh": "📌 架构与实施要点 (NOTE)",
        "title_ja": "📌 アーキテクチャ・設計要点 (NOTE)",
        "title_color": RGBColor(37, 99, 235),
    },
    "IMPORTANT": {
        "border_color": "7C3AED",
        "bg_color": "F5F3FF",
        "title_zh": "⚡ 重要关键规则 (IMPORTANT)",
        "title_ja": "⚡ 重要確認事項 (IMPORTANT)",
        "title_color": RGBColor(124, 58, 237),
    },
    "WARNING": {
        "border_color": "D97706",
        "bg_color": "FFFBEB",
        "title_zh": "⚠️ 生产防坑警告 (WARNING)",
        "title_ja": "⚠️ 運用上の注意事項・警告 (WARNING)",
        "title_color": RGBColor(217, 119, 6),
    },
    "CAUTION": {
        "border_color": "DC2626",
        "bg_color": "FEF2F2",
        "title_zh": "🚨 金融级安全军规 (CAUTION)",
        "title_ja": "🚨 金融セキュリティ遵守事項 (CAUTION)",
        "title_color": RGBColor(220, 38, 38),
    },
}

def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    cell._tc.get_or_add_tcPr().append(tcMar)

def set_cell_borders(cell, top="none", bottom="none", left="none", right="none", 
                     sz="4", color="CBD5E1"):
    borders_xml = f'<w:tcBorders {nsdecls("w")}>'
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        if val == "single":
            borders_xml += f'<w:{side} w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        elif val == "thick":
            borders_xml += f'<w:{side} w:val="single" w:sz="32" w:space="0" w:color="{color}"/>'
        else:
            borders_xml += f'<w:{side} w:val="none"/>'
    borders_xml += '</w:tcBorders>'
    cell._tc.get_or_add_tcPr().append(parse_xml(borders_xml))

def add_page_number(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_total_pages(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> NUMPAGES </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def apply_inline_formatting(p, text, default_font="Microsoft YaHei", default_size=10.5, default_color=COLOR_TEXT_MAIN):
    """
    Parses **bold**, `code`, *italic*, and [link](url)
    """
    pattern = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|\[[^\]]+\]\([^)]+\))')
    tokens = pattern.split(text)

    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = p.add_run(token[2:-2])
            run.font.name = default_font
            run.font.size = Pt(default_size)
            run.font.bold = True
            run.font.color.rgb = default_color
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            code_text = token[1:-1]
            run = p.add_run(f" {code_text} ")
            run.font.name = "Consolas"
            run.font.size = Pt(max(default_size - 1, 9.0))
            run.font.color.rgb = RGBColor(185, 28, 28) # Tech red-brick
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.font.name = default_font
            run.font.size = Pt(default_size)
            run.font.italic = True
            run.font.color.rgb = default_color
        elif token.startswith('[') and '](' in token and token.endswith(')'):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', token)
            if m:
                link_text, url = m.groups()
                run = p.add_run(link_text)
                run.font.name = default_font
                run.font.size = Pt(default_size)
                run.font.color.rgb = COLOR_SECONDARY
                run.font.underline = True
            else:
                run = p.add_run(token)
                run.font.name = default_font
                run.font.size = Pt(default_size)
                run.font.color.rgb = default_color
        else:
            # Clean HTML tags like <br>, <strong> if any
            clean = token.replace('<br>', '\n').replace('<br/>', '\n')
            clean = re.sub(r'</?[a-zA-Z0-9]+[^>]*>', '', clean)
            run = p.add_run(clean)
            run.font.name = default_font
            run.font.size = Pt(default_size)
            run.font.color.rgb = default_color

def create_cover_page(doc, lang="zh"):
    """
    Builds a high-end corporate cover page.
    """
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.different_first_page_header_footer = True

    # Header / Footer for subsequent pages
    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if lang == "zh":
        hrun = header_p.add_run("银行证券 DX 实战项目：Oracle 11g 至 Snowflake 迁移架构与实施手册")
    else:
        hrun = header_p.add_run("銀行証券 DX 実践プロジェクト：Oracle 11g から Snowflake への企業級データ移行実務")
    hrun.font.name = "Segoe UI"
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = RGBColor(148, 163, 184)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if lang == "zh":
        frun1 = footer_p.add_run("【内部资料 · 严格保密】                                                    第 ")
    else:
        frun1 = footer_p.add_run("【社外秘・Confidential】                                                    ページ ")
    frun1.font.name = "Segoe UI"
    frun1.font.size = Pt(8.5)
    frun1.font.color.rgb = RGBColor(148, 163, 184)
    add_page_number(frun1)
    
    frun2 = footer_p.add_run(" / 共 ")
    frun2.font.name = "Segoe UI"
    frun2.font.size = Pt(8.5)
    frun2.font.color.rgb = RGBColor(148, 163, 184)
    add_total_pages(frun2)
    frun3 = footer_p.add_run(" 页" if lang == "zh" else " 頁")
    frun3.font.name = "Segoe UI"
    frun3.font.size = Pt(8.5)
    frun3.font.color.rgb = RGBColor(148, 163, 184)

    # Top spacing
    for _ in range(3):
        doc.add_paragraph()

    # Category Badge
    badge_p = doc.add_paragraph()
    badge_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    badge_run = badge_p.add_run("ENTERPRISE CLOUD DATA MIGRATION BLUEPRINT" if lang == "zh" else "ENTERPRISE CLOUD DATA MIGRATION BLUEPRINT")
    badge_run.font.name = "Segoe UI"
    badge_run.font.size = Pt(10)
    badge_run.font.bold = True
    badge_run.font.color.rgb = COLOR_SECONDARY

    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(16)
    if lang == "zh":
        title_run = title_p.add_run("银行证券 DX 实战项目\nOracle 11g 迁移至 Snowflake\n企业级落地全景指南")
    else:
        title_run = title_p.add_run("銀行証券 DX 実践プロジェクト\nOracle 11g から Snowflake への\n企業級データ移行実務マニュアル")
    title_run.font.name = "Microsoft YaHei" if lang == "zh" else "Meiryo"
    title_run.font.size = Pt(26)
    title_run.font.bold = True
    title_run.font.color.rgb = COLOR_PRIMARY

    # Accent decorative bar (Table)
    bar_tbl = doc.add_table(rows=1, cols=1)
    bar_tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    bar_cell = bar_tbl.cell(0, 0)
    bar_cell.width = Inches(6.2)
    set_cell_background(bar_cell, "2563EB")
    set_cell_margins(bar_cell, top=20, bottom=20, left=0, right=0)
    set_cell_borders(bar_cell, "none", "none", "none", "none")
    bp = bar_cell.paragraphs[0]
    bp.paragraph_format.space_before = Pt(0)
    bp.paragraph_format.space_after = Pt(0)

    # Subtitle
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(18)
    sub_p.paragraph_format.space_after = Pt(40)
    if lang == "zh":
        sub_text = "金融级 3 层 RBAC 权限隔离 · 高性能 COPY INTO · 4 阶段新旧数据突合对账全真还原"
    else:
        sub_text = "金融機関向け 3 層 RBAC 権限分離 · 高速 COPY INTO · 4 段階新旧データ突合検証の完全再現"
    sub_run = sub_p.add_run(sub_text)
    sub_run.font.name = "Microsoft YaHei" if lang == "zh" else "Meiryo"
    sub_run.font.size = Pt(12)
    sub_run.font.color.rgb = COLOR_TEXT_MUTED

    # Middle spacing
    for _ in range(4):
        doc.add_paragraph()

    # Metadata Box
    meta_tbl = doc.add_table(rows=5, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_tbl.autofit = False

    meta_items_zh = [
        ("文档版本", "v1.0.0 (Release Edition)"),
        ("适用范围", "银行证券 核心交易明细数仓迁移 (CUST_TRADE_DETAIL)"),
        ("编制机构", "银行证券 DX 迁移项目推进委员会 & 云原生架构组"),
        ("发布日期", "2026 年 10 月"),
        ("密级分类", "【内部资料 · 严格保密 · 未经授权禁止扩散】"),
    ]
    meta_items_ja = [
        ("ドキュメント版数", "v1.0.0 (正式リリース版)"),
        ("適用範囲", "銀行証券 コア取引明細 DWH 移行実務 (CUST_TRADE_DETAIL)"),
        ("作成組織", "銀行証券 DX 移行プロジェクト推進室 & クラウド推進チーム"),
        ("発行年月", "2026 年 10 月"),
        ("機密区分", "【社外秘・Confidential・第三者開示禁止】"),
    ]
    items = meta_items_zh if lang == "zh" else meta_items_ja

    for r_idx, (label, val) in enumerate(items):
        c0 = meta_tbl.cell(r_idx, 0)
        c1 = meta_tbl.cell(r_idx, 1)
        c0.width = Inches(1.8)
        c1.width = Inches(4.4)

        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=90, bottom=90, left=140, right=140)
        set_cell_margins(c1, top=90, bottom=90, left=140, right=140)
        set_cell_borders(c0, top="single", bottom="single", left="none", right="none", color="E2E8F0")
        set_cell_borders(c1, top="single", bottom="single", left="none", right="none", color="E2E8F0")

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = "Microsoft YaHei" if lang == "zh" else "Meiryo"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_TEXT_MUTED

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = "Microsoft YaHei" if lang == "zh" else "Meiryo"
        r1.font.size = Pt(9.5)
        if r_idx == 4:
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(220, 38, 38)
        else:
            r1.font.color.rgb = COLOR_TEXT_MAIN

    doc.add_page_break()

def create_table_of_contents_page(doc, lang="zh"):
    """
    Creates an executive TOC summary table and overview.
    """
    h_p = doc.add_paragraph()
    h_p.paragraph_format.space_before = Pt(12)
    h_p.paragraph_format.space_after = Pt(18)
    h_run = h_p.add_run("目录与全书架构体系" if lang == "zh" else "目次とプロジェクト全体体系")
    h_run.font.name = "Microsoft YaHei" if lang == "zh" else "Meiryo"
    h_run.font.size = Pt(18)
    h_run.font.bold = True
    h_run.font.color.rgb = COLOR_PRIMARY

    intro_p = doc.add_paragraph()
    intro_p.paragraph_format.space_after = Pt(16)
    if lang == "zh":
        intro_text = "本项目以真实证券交易明细为载体，涵盖从架构顶层设计、账号体系规范、Oracle源端数据建模，到 Snowflake 生产级 3 层 RBAC 权限管控、高效数据抽取加密入库及 4 阶段全严密突合对账等全周期环节："
    else:
        intro_text = "本プロジェクトは、証券取引の顧客取引明細データを題材に、アーキテクチャ設計から 3 層 RBAC 権限構築、Oracle 移行元データ準備、COPY INTO 取込実務、そして 4 段階の突合監査に至る全移行工程を完全網羅しています："
    apply_inline_formatting(intro_p, intro_text, "Microsoft YaHei" if lang == "zh" else "Meiryo", 10.5)

    toc_data_zh = [
        ("第 00 章", "业务背景、总体架构设计与成本测算模型", "架构设计 · 存储计算分离 · 计费仿真 ($508/月) · 8大类型映射"),
        ("第 01 章", "Snowflake 账号申请、初始化与控制台本地化", "Enterprise 试用申请 · AWS 东京区域合规 · 日语控制台切换 · 预置角色"),
        ("第 02 章", "Oracle 11g 迁移源环境搭建与模拟数据准备", "Docker 部署 · VSCode 数据库连接 · 核心表 DDL · 11 组高覆盖率测试数据"),
        ("第 03 章", "Snowflake 金融级 3 层 RBAC 权限模型实战", "USERADMIN 与 SYSADMIN 职责分离 · 3 层角色矩阵 (ADMIN/ENG/ANALYST)"),
        ("第 04 章", "数据迁移实操：Stage 暂存区与 COPY INTO 加载", "CSV 加密导出 · SnowSQL PUT 上传 · 批量类型转换 · 执行统计分析"),
        ("第 05 章", "两方数据严密突合对账与一致性稽核方案", "4 阶段对账体系：行数对账 · 金额分项对账 · HASH 指纹比对 · 差分提取"),
        ("第 06 章", "生产级避坑指南、排障实录与运维最佳实践", "常见报错分析 · 算力成本优化 · 时区规避 · 数据类型溢出防范"),
    ]

    toc_data_ja = [
        ("第 00 章", "プロジェクト背景・アーキテクチャ設計・コスト試算", "アーキテクチャ全体像 · 課金体系 ($508/月モデル) · 8大データ型設計方針"),
        ("第 01 章", "Snowflake アカウント開設と初期設定", "Enterprise 試用申請 · AWS 東京リージョン選定 · 日本語 UI 設定 · 組み込みロール"),
        ("第 02 章", "Oracle 移行元環境構築と模擬データ投入", "Docker 11g 構築 · VSCode DB 接続 · テーブル DDL · 11 件の極限境界値データ"),
        ("第 03 章", "金融機関向け 3 層 RBAC 権限モデルと環境構築", "USERADMIN と SYSADMIN の職責分離 · 3層ロール設計 (ADMIN/ENG/ANALYST)"),
        ("第 04 章", "データ移行実務（抽出・ステージ・取込）", "CSV 抽出 · SnowSQL PUT アップロード · COPY INTO 型変換 · ロード検証"),
        ("第 05 章", "移行後データ突合検証・整合性監査", "4 段階突合検証：総件数突合 · 金額集計突合 · HASH 指紋検証 · 差分特定"),
        ("第 06 章", "トラブルシューティングとベストプラクティス", "エラー逆引き解決 · コンピュートコスト最適化 · タイムゾーン対策 · 溢れ防止"),
    ]

    toc_items = toc_data_zh if lang == "zh" else toc_data_ja

    tbl = doc.add_table(rows=len(toc_items) + 1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    # Header
    headers = ["章节编号", "章节核心主题", "涵盖关键技术要点"] if lang == "zh" else ["章番号", "主要テーマ", "網羅する重要技術要素"]
    col_widths = [Inches(1.1), Inches(2.5), Inches(2.6)]

    for c_idx, h_text in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        cell.width = col_widths[c_idx]
        set_cell_background(cell, COLOR_HEADER_BG)
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        set_cell_borders(cell, top="none", bottom="single", left="none", right="none", color="1E3A8A")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_text)
        r.font.name = "Microsoft YaHei" if lang == "zh" else "Meiryo"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, (ch_num, ch_title, ch_desc) in enumerate(toc_items, start=1):
        bg = COLOR_ZEBRA_BG if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, text_val in enumerate([ch_num, ch_title, ch_desc]):
            cell = tbl.cell(r_idx, c_idx)
            cell.width = col_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            set_cell_borders(cell, top="single", bottom="single", left="none", right="none", color="E2E8F0")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text_val)
            r.font.name = "Microsoft YaHei" if lang == "zh" else "Meiryo"
            r.font.size = Pt(9.0)
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = COLOR_SECONDARY
            elif c_idx == 1:
                r.font.bold = True
                r.font.color.rgb = COLOR_PRIMARY
            else:
                r.font.color.rgb = COLOR_TEXT_MUTED

    doc.add_page_break()

def parse_and_append_markdown(doc, md_path, lang="zh"):
    """
    Parses a single Markdown chapter and converts it into high-fidelity Docx elements.
    """
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    i = 0
    num_lines = len(lines)

    in_code_block = False
    code_lang = ""
    code_lines = []

    in_table = False
    table_lines = []

    in_callout = False
    callout_type = "NOTE"
    callout_lines = []

    font_main = "Microsoft YaHei" if lang == "zh" else "Meiryo"

    def flush_code():
        nonlocal in_code_block, code_lang, code_lines
        if not code_lines:
            in_code_block = False
            return
        
        raw_code = '\n'.join(code_lines)
        code_lines = []
        in_code_block = False

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.2)
        set_cell_background(cell, COLOR_CODE_BG)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        set_cell_borders(cell, top="single", bottom="single", left="single", right="single", sz="4", color="CBD5E1")

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15

        if code_lang:
            badge = p.add_run(f"[{code_lang.upper()}]\n")
            badge.font.name = "Consolas"
            badge.font.size = Pt(8.5)
            badge.font.bold = True
            badge.font.color.rgb = COLOR_SECONDARY

        c_run = p.add_run(raw_code)
        c_run.font.name = "Consolas"
        c_run.font.size = Pt(8.5)
        c_run.font.color.rgb = COLOR_CODE_TEXT

        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(4)

    def flush_callout():
        nonlocal in_callout, callout_type, callout_lines
        if not callout_lines:
            in_callout = False
            return

        cfg = CALLOUT_CONFIGS.get(callout_type, CALLOUT_CONFIGS["NOTE"])
        full_text = '\n'.join(callout_lines).strip()
        callout_lines = []
        in_callout = False

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.2)
        set_cell_background(cell, cfg["bg_color"])
        set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
        set_cell_borders(cell, top="none", bottom="none", left="thick", right="none", color=cfg["border_color"])

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25

        title_str = cfg["title_zh"] if lang == "zh" else cfg["title_ja"]
        t_run = p.add_run(f"{title_str}\n")
        t_run.font.name = font_main
        t_run.font.size = Pt(10)
        t_run.font.bold = True
        t_run.font.color.rgb = cfg["title_color"]

        apply_inline_formatting(p, full_text, font_main, 9.5, COLOR_TEXT_MAIN)

        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(4)

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            in_table = False
            return

        raw_rows = []
        for line in table_lines:
            parts = [p.strip() for p in line.split('|')[1:-1]]
            if parts:
                raw_rows.append(parts)
        table_lines = []
        in_table = False

        if len(raw_rows) < 2:
            return

        # Check if row 1 is delimiter (| :--- | :--- |)
        headers = raw_rows[0]
        data_rows = []
        for row in raw_rows[1:]:
            if all(re.match(r'^:?-+:?$', c) for c in row):
                continue
            data_rows.append(row)

        num_cols = len(headers)
        if num_cols == 0:
            return

        tbl = doc.add_table(rows=len(data_rows) + 1, cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        total_width = 6.2
        col_w = Inches(total_width / num_cols)

        # Header Row
        for c_idx, h_text in enumerate(headers):
            cell = tbl.cell(0, c_idx)
            cell.width = col_w
            set_cell_background(cell, COLOR_HEADER_BG)
            set_cell_margins(cell, top=130, bottom=130, left=130, right=130)
            set_cell_borders(cell, top="none", bottom="single", left="none", right="none", color="1E3A8A")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            apply_inline_formatting(p, h_text, font_main, 9.0, RGBColor(255, 255, 255))
            for r in p.runs:
                r.font.bold = True

        # Data Rows
        for r_idx, row in enumerate(data_rows, start=1):
            bg = COLOR_ZEBRA_BG if r_idx % 2 == 1 else "FFFFFF"
            for c_idx in range(num_cols):
                cell = tbl.cell(r_idx, c_idx)
                cell.width = col_w
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=100, bottom=100, left=130, right=130)
                set_cell_borders(cell, top="single", bottom="single", left="none", right="none", color="E2E8F0")
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                c_val = row[c_idx] if c_idx < len(row) else ""
                apply_inline_formatting(p, c_val, font_main, 8.5, COLOR_TEXT_MAIN)

        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(4)

    while i < num_lines:
        line = lines[i]

        # 1. Code blocks
        if line.strip().startswith('```'):
            if in_code_block:
                flush_code()
            else:
                if in_callout:
                    flush_callout()
                if in_table:
                    flush_table()
                in_code_block = True
                code_lang = line.strip()[3:].strip()
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # 2. Callouts / Blockquotes
        if line.strip().startswith('>'):
            if in_table:
                flush_table()
            stripped = line.strip()[1:].strip()
            # Match alert header > [!TIP]
            m_alert = re.match(r'^\[!(TIP|NOTE|IMPORTANT|WARNING|CAUTION)\]', stripped, re.IGNORECASE)
            if m_alert:
                if in_callout:
                    flush_callout()
                in_callout = True
                callout_type = m_alert.group(1).upper()
                callout_lines = []
            else:
                if not in_callout:
                    in_callout = True
                    callout_type = "NOTE"
                    callout_lines = []
                callout_lines.append(stripped)
            i += 1
            continue
        elif in_callout and (line.strip() == '' or not line.strip().startswith('>')):
            flush_callout()

        # 3. Markdown Tables
        if line.strip().startswith('|') and line.strip().endswith('|'):
            in_table = True
            table_lines.append(line.strip())
            i += 1
            continue
        elif in_table:
            flush_table()

        stripped_line = line.strip()

        # Empty lines
        if not stripped_line:
            i += 1
            continue

        # Frontmatter skip if any
        if stripped_line == '---':
            i += 1
            continue

        # Headings
        if stripped_line.startswith('# ') and not stripped_line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(20)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.keep_with_next = True
            title_text = stripped_line[2:].strip()
            r = p.add_run(title_text)
            r.font.name = font_main
            r.font.size = Pt(20)
            r.font.bold = True
            r.font.color.rgb = COLOR_PRIMARY
            i += 1
            continue

        if stripped_line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            title_text = stripped_line[3:].strip()
            r = p.add_run(title_text)
            r.font.name = font_main
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = COLOR_PRIMARY
            i += 1
            continue

        if stripped_line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            title_text = stripped_line[4:].strip()
            r = p.add_run(title_text)
            r.font.name = font_main
            r.font.size = Pt(11.5)
            r.font.bold = True
            r.font.color.rgb = COLOR_SECONDARY
            i += 1
            continue

        if stripped_line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            title_text = stripped_line[5:].strip()
            r = p.add_run(title_text)
            r.font.name = font_main
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = COLOR_TEXT_MAIN
            i += 1
            continue

        # Images: ![alt](path)
        img_match = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', stripped_line)
        if img_match:
            alt_text, img_rel_path = img_match.groups()
            # Resolve image path relative to md file
            md_dir = md_path.parent
            img_abs_path = (md_dir / img_rel_path).resolve()
            if img_abs_path.exists():
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(10)
                p_img.paragraph_format.space_after = Pt(4)
                doc_pic = p_img.add_run()
                doc_pic.add_picture(str(img_abs_path), width=Inches(5.6))

                # Image Caption
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(2)
                p_cap.paragraph_format.space_after = Pt(12)
                cap_run = p_cap.add_run(f"▲ {alt_text}" if alt_text else "▲ 参考図・スクリーンショット")
                cap_run.font.name = font_main
                cap_run.font.size = Pt(8.5)
                cap_run.font.italic = True
                cap_run.font.color.rgb = COLOR_TEXT_MUTED
            i += 1
            continue

        # Lists: Unordered (- or *)
        if re.match(r'^[\*\-]\s+', stripped_line):
            list_text = re.sub(r'^[\*\-]\s+', '', stripped_line)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            bullet_run = p.add_run("•  ")
            bullet_run.font.name = "Arial"
            bullet_run.font.bold = True
            bullet_run.font.color.rgb = COLOR_SECONDARY
            apply_inline_formatting(p, list_text, font_main, 10.0, COLOR_TEXT_MAIN)
            i += 1
            continue

        # Lists: Ordered (1. )
        m_num = re.match(r'^(\d+)\.\s+(.*)', stripped_line)
        if m_num:
            num_str, list_text = m_num.groups()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            n_run = p.add_run(f"{num_str}. ")
            n_run.font.name = font_main
            n_run.font.bold = True
            n_run.font.color.rgb = COLOR_SECONDARY
            apply_inline_formatting(p, list_text, font_main, 10.0, COLOR_TEXT_MAIN)
            i += 1
            continue

        # Normal Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.3
        apply_inline_formatting(p, stripped_line, font_main, 10.0, COLOR_TEXT_MAIN)
        i += 1

    # Cleanup at EOF
    if in_code_block:
        flush_code()
    if in_callout:
        flush_callout()
    if in_table:
        flush_table()

def generate_full_document(lang="zh"):
    doc = docx.Document()
    
    # Page setup
    for sec in doc.sections:
        sec.top_margin = Inches(0.9)
        sec.bottom_margin = Inches(0.9)
        sec.left_margin = Inches(0.9)
        sec.right_margin = Inches(0.9)

    print(f"[{lang.upper()}] Creating Cover Page...")
    create_cover_page(doc, lang)

    print(f"[{lang.upper()}] Creating Table of Contents...")
    create_table_of_contents_page(doc, lang)

    source_dir = BASE_DIR / "docs" / lang
    for filename, ch_id in CHAPTERS:
        file_path = source_dir / filename
        if not file_path.exists():
            print(f"Warning: {file_path} not found.")
            continue
        print(f"[{lang.upper()}] Processing Chapter {ch_id}: {filename}...")
        parse_and_append_markdown(doc, file_path, lang)
        # Add page break between chapters (except last one)
        if ch_id != CHAPTERS[-1][1]:
            doc.add_page_break()

    if lang == "zh":
        out_filename = "银行证券_Oracle至Snowflake数据迁移全套实施方案(中文版).docx"
    else:
        out_filename = "銀行証券_OracleからSnowflakeへのデータ移行実務マニュアル(日本語版).docx"

    out_path_exports = EXPORTS_DIR / out_filename
    out_path_root = BASE_DIR / out_filename

    doc.save(str(out_path_exports))
    doc.save(str(out_path_root))
    print(f"[{lang.upper()}] Successfully saved to:")
    print(f"  - {out_path_root}")
    print(f"  - {out_path_exports}")

if __name__ == "__main__":
    print("=== Generating Enterprise Word Documents (.docx) ===")
    generate_full_document("zh")
    generate_full_document("ja")
    print("=== All Documents Generated Successfully ===")
